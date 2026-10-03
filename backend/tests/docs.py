"""Builds small PDF and DOCX files in memory, so tests need no binary fixtures."""

import io

import docx
import pymupdf


def make_pdf(pages: list[list[tuple[str, float]]], header: str | None = None, footer: bool = False) -> bytes:
    """pages: per page, a list of (text, font size). Larger sizes become headings."""
    doc = pymupdf.open()
    for number, blocks in enumerate(pages, start=1):
        page = doc.new_page()
        y = 50.0
        if header:
            page.insert_text((72, 30), header, fontsize=9)
        for text, size in blocks:
            box = pymupdf.Rect(72, y, 540, y + 400)
            used = page.insert_textbox(box, text, fontsize=size)
            y += (400 - used) + size * 1.5
        if footer:
            page.insert_text((300, 800), str(number), fontsize=9)
    data = doc.tobytes()
    doc.close()
    return data


def make_docx() -> bytes:
    document = docx.Document()
    document.add_heading("Cell Biology", level=1)
    document.add_paragraph("Cells are the basic unit of life. Every living thing is made of cells.")
    document.add_heading("Organelles", level=2)
    document.add_paragraph("The nucleus holds DNA.", style="List Bullet")
    document.add_paragraph("Mitochondria release energy.", style="List Bullet")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Organelle"
    table.cell(0, 1).text = "Job"
    table.cell(1, 0).text = "Ribosome"
    table.cell(1, 1).text = "Makes proteins"
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def make_png(size: tuple[int, int] = (400, 300), colour: str = "white") -> bytes:
    from PIL import Image

    buffer = io.BytesIO()
    Image.new("RGB", size, colour).save(buffer, format="PNG")
    return buffer.getvalue()


def make_scanned_pdf(text_pages: list[str | None]) -> bytes:
    """Per page: real text, or None for a page that is only a picture (a scan)."""
    doc = pymupdf.open()
    picture = make_png((600, 800), "#f4f1e8")
    for text in text_pages:
        page = doc.new_page()
        if text is None:
            page.insert_image(page.rect, stream=picture)
        else:
            page.insert_textbox(pymupdf.Rect(72, 72, 540, 700), text, fontsize=11)
    data = doc.tobytes()
    doc.close()
    return data


def make_video(path, scenes: list[tuple[str, int]], audio: bool = True) -> None:
    """A small MP4: one solid colour per scene (colour, seconds), with a tone as the soundtrack."""
    import subprocess

    inputs: list[str] = []
    for colour, seconds in scenes:
        inputs += ["-f", "lavfi", "-i", f"color=c={colour}:s=320x180:r=10:d={seconds}"]
    total = sum(s for _, s in scenes)
    concat = "".join(f"[{i}:v]" for i in range(len(scenes))) + f"concat=n={len(scenes)}:v=1:a=0[v]"
    cmd = ["ffmpeg", "-nostdin", "-v", "error", "-y", *inputs]
    if audio:
        cmd += ["-f", "lavfi", "-i", f"sine=frequency=440:duration={total}"]
    cmd += ["-filter_complex", concat, "-map", "[v]"]
    if audio:
        cmd += ["-map", f"{len(scenes)}:a", "-c:a", "aac"]
    cmd += ["-c:v", "libx264", "-g", "20", "-pix_fmt", "yuv420p", str(path)]
    subprocess.run(cmd, check=True, capture_output=True)


def make_audio(path, seconds: int) -> None:
    import subprocess

    subprocess.run(
        [
            "ffmpeg",
            "-nostdin",
            "-v",
            "error",
            "-y",
            "-f",
            "lavfi",
            "-i",
            f"sine=frequency=300:duration={seconds}",
            "-c:a",
            "libmp3lame",
            str(path),
        ],
        check=True,
        capture_output=True,
    )
