"""Exports render the same guide, with the same labels, safely, in every format."""

import io

import docx
import pymupdf
import pytest

from app.export import FORMATS, ExportInfo, filename, render
from app.export.common import latex_to_text, math_to_text
from app.pipeline import orchestrator
from tests.fakes import FakeLLM
from tests.pipeline.test_orchestrator import _ctx

INFO = ExportInfo(title="Photosynthesis", level="high_school", source_titles=["Biology chapter 8"])


@pytest.fixture(scope="module")
def content():
    text = (__import__("pathlib").Path(__file__).parents[1] / "fixtures" / "photosynthesis.md").read_text()
    guide = orchestrator.run(_ctx(text, FakeLLM(review_scores={"Chlorophyll": [40, 45]})))
    first = guide["concepts"][0]
    first["concept"] += (
        "\n\nThe equation is $6CO_2 + 6H_2O \\rightarrow C_6H_{12}O_6 + 6O_2$ and **light** matters.\n\n"
        "<script>alert(1)</script> ![tracker](http://169.254.169.254/latest) [link](javascript:alert(1))"
    )
    return guide


def _pdf_text(data: bytes) -> str:
    return " ".join(page.get_text() for page in pymupdf.open(stream=data, filetype="pdf"))


def _docx_text(data: bytes) -> str:
    document = docx.Document(io.BytesIO(data))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        parts += [cell.text for row in table.rows for cell in row.cells]
    return "\n".join(parts)


@pytest.mark.parametrize("fmt", list(FORMATS))
def test_every_format_has_the_whole_guide_and_its_labels(content, fmt):
    data = render(content, fmt, INFO)
    if fmt == "pdf":
        assert data.startswith(b"%PDF")
        text = _pdf_text(data)
    elif fmt == "docx":
        text = _docx_text(data)
    else:
        text = data.decode()
    for expected in (
        "Photosynthesis",
        "What is this topic about?",
        "Key vocabulary",
        "Inputs and outputs",
        "Light reactions",
        "Explain like I'm new",
        "AI-generated example",
        "From your material",
        "Check this against your source",
        "says something the source does not say",
        "Practice questions",
        "Answer key",
        "Flashcards",
        "Review checklist",
        "Biology chapter 8",
        "part 1",
    ):
        assert expected in text, f"{fmt} is missing {expected!r}"


def test_html_is_safe(content):
    html = render(content, "html", INFO).decode()
    assert "<script>" not in html and "&lt;script&gt;" in html
    assert "<img" not in html and "href=" not in html
    assert "<svg" in html  # the diagram is drawn inline


def test_pdf_and_word_show_maths_as_text_and_draw_the_diagram(content):
    assert "6CO₂ + 6H₂O → C₆H₁₂O₆ + 6O₂" in _pdf_text(render(content, "pdf", INFO))
    data = render(content, "docx", INFO)
    assert "6CO₂ + 6H₂O → C₆H₁₂O₆ + 6O₂" in _docx_text(data)
    assert len(docx.Document(io.BytesIO(data)).inline_shapes) == 1


def test_markdown_keeps_maths_and_mermaid_for_markdown_apps(content):
    md = render(content, "md", INFO).decode()
    assert "$6CO_2" in md and "```mermaid\nflowchart LR" in md and "Diagram as text" in md
    assert "- [ ] I can explain photosynthesis." in md


def test_latex_to_text():
    assert latex_to_text(r"v = \frac{\Delta x}{\Delta t}") == "v = Δx/Δt"
    assert latex_to_text(r"x = \frac{-b \pm \sqrt{b^2-4ac}}{2a}") == "x = (-b ± √(b²-4ac))/2a"
    assert latex_to_text(r"K_{eq}") == "K_(eq)"
    assert math_to_text("It costs $5 and $10, and $E = mc^2$.") == "It costs $5 and $10, and E = mc²."


def test_filename():
    assert filename("Photosynthesis: Light & Dark!", "pdf") == "photosynthesis-light-dark.pdf"
    assert filename("???", "md") == "study-guide.md"
