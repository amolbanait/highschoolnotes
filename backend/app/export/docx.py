"""The guide as a Word document (python-docx).

Model-written Markdown becomes real Word formatting (bold, italics, lists). Diagrams are drawn
with the same SVG renderer as the PDF and embedded as pictures; formulas are readable text.
"""

import io

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from markdown_it import MarkdownIt

from app.diagrams import mermaid, svg
from app.export.common import (
    AI_ANALOGY,
    AI_EXAMPLE,
    CHECK_THIS,
    CHECK_THIS_BODY,
    DIFFICULTY_LABELS,
    FACT_KINDS,
    FOOTER,
    FROM_SOURCE,
    LEVEL_TABS,
    QUESTION_TYPES,
    UNVERIFIED,
    ExportInfo,
    concept_titles,
    latex_to_text,
    math_to_text,
)

_md = MarkdownIt("commonmark", {"html": False}).enable("table")
_md.disable(["image", "link", "autolink"])  # model text never links or loads anything

GREY = RGBColor(0x64, 0x74, 0x8B)
GREEN = RGBColor(0x06, 0x5F, 0x46)
AMBER = RGBColor(0x92, 0x40, 0x0E)
INDIGO = RGBColor(0x43, 0x38, 0xCA)


def to_docx(content: dict, info: ExportInfo) -> bytes:
    doc = Document()
    _setup(doc, content.get("title") or info.title)
    titles = concept_titles(content)

    def cite(paragraph, refs) -> None:
        label = info.cite(refs)
        if label:
            _note(paragraph, " " + label, GREY)

    doc.add_heading(content.get("title") or info.title, level=0)
    meta = [f"Level: {info.level_label}"]
    if info.source_titles:
        meta.append("From: " + "; ".join(info.source_titles))
    meta.append(f"Made {info.made_on_label}")
    _note(doc.add_paragraph(), " · ".join(meta), GREY)

    if content.get("warnings"):
        cell = _callout(doc, "FFFBEB")
        _bold_line(cell.paragraphs[0], "Double-check these")
        for x in content["warnings"]:
            cell.add_paragraph(x, style="List Bullet")

    if content.get("overview"):
        doc.add_heading("What is this topic about?", level=1)
        _blocks(doc, content["overview"])
    if content.get("objectives"):
        _bold_line(doc.add_paragraph(), "By the end, you should be able to:")
        for o in content["objectives"]:
            _inline(doc.add_paragraph(style="List Bullet"), o)

    if content.get("vocabulary"):
        doc.add_heading("Key vocabulary", level=1)
        table = _table(doc, ["Term", "Simple meaning", "Example"])
        for v in content["vocabulary"]:
            row = table.add_row().cells
            _bold_line(row[0].paragraphs[0], v["term"])
            _inline(row[1].paragraphs[0], v["definition"])
            cite(row[1].paragraphs[0], v.get("source_refs"))
            if v.get("verified") is False:
                _note(row[1].add_paragraph(), "⚠ " + UNVERIFIED, AMBER)
            _inline(row[2].paragraphs[0], v.get("example") or "")

    concepts = content.get("concepts", [])
    if concepts:
        doc.add_heading("Main concepts", level=1)
    for number, c in enumerate(concepts, start=1):
        _concept(doc, c, number, titles, cite)

    if content.get("facts"):
        doc.add_heading("Key facts to remember", level=1)
        for f in content["facts"]:
            p = doc.add_paragraph(style="List Bullet")
            tag = FACT_KINDS.get(f["kind"], "Fact") + (
                ", teacher emphasis" if f.get("teacher_emphasis") else ""
            )
            _bold_line(p, f"{tag}: ")
            _inline(p, f["text"])
            if f.get("quote"):
                q = doc.add_paragraph()
                q.paragraph_format.left_indent = Inches(0.4)
                run = q.add_run(f"“{f['quote']}”")
                run.italic = True
                cite(q, f.get("source_refs"))
            if f.get("verified") is False:
                _note(_indented(doc), "⚠ " + UNVERIFIED, AMBER)

    if content.get("formulas"):
        doc.add_heading("Formulas", level=1)
        for f in content["formulas"]:
            p = doc.add_paragraph()
            run = p.add_run(latex_to_text(f["latex"]))
            run.font.size = Pt(13)
            run.font.name = "Cambria Math"
            _blocks(doc, f["meaning"])
            cite(doc.paragraphs[-1], f.get("source_refs"))
            if f.get("verified") is False:
                _note(doc.add_paragraph(), "⚠ " + UNVERIFIED, AMBER)

    if content.get("relationships"):
        doc.add_heading("How the ideas connect", level=1)
        for r in content["relationships"]:
            p = doc.add_paragraph(style="List Bullet")
            _bold_line(p, titles.get(r["from"], r["from"]))
            p.add_run(f" → {r['type']} → ")
            _bold_line(p, titles.get(r["to"], r["to"]))
            p.add_run(": ")
            _inline(p, r["explanation"])

    if content.get("summary"):
        doc.add_heading("Summary", level=1)
        _blocks(doc, content["summary"])

    questions = content.get("questions", [])
    if questions:
        doc.add_heading("Practice questions", level=1)
        for i, q in enumerate(questions, start=1):
            p = doc.add_paragraph()
            _bold_line(p, f"{i}. ")
            _note(p, f"({QUESTION_TYPES.get(q['type'], q['type'])}) ", GREY)
            _inline(p, q["prompt"])
            for j, option in enumerate(q.get("options") or []):
                o = _indented(doc)
                o.add_run(f"{chr(65 + j)}. ")
                _inline(o, option)

    if content.get("review_checklist"):
        doc.add_heading("Review checklist", level=1)
        doc.add_paragraph("Tick each one you could explain to a friend without notes.")
        for item in content["review_checklist"]:
            p = doc.add_paragraph()
            p.add_run("☐  ")
            _inline(p, item)

    if content.get("not_covered"):
        doc.add_heading("Not clearly covered in your material", level=1)
        doc.add_paragraph("You might expect these here. Check your textbook or ask your teacher.")
        for n in content["not_covered"]:
            doc.add_paragraph(n, style="List Bullet")

    if questions:
        doc.add_page_break()
        doc.add_heading("Answer key", level=1)
        for i, q in enumerate(questions, start=1):
            answer = q["answer"]
            if q.get("options") and answer in q["options"]:
                answer = f"{chr(65 + q['options'].index(answer))}. {answer}"
            p = doc.add_paragraph()
            _bold_line(p, f"{i}. ")
            _bold_line(p, math_to_text(answer))
            if q.get("explanation"):
                p.add_run(" ")
                _inline(p, " ".join(q["explanation"].split()))
            cite(p, q.get("source_refs"))

    if content.get("flashcards"):
        doc.add_heading("Flashcards", level=1)
        table = _table(doc, ["Front", "Back"])
        for card in content["flashcards"]:
            row = table.add_row().cells
            _inline(row[0].paragraphs[0], card["front"])
            _inline(row[1].paragraphs[0], card["back"])

    _note(doc.add_paragraph(), FOOTER, GREY)
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def _concept(doc, c: dict, number: int, titles: dict[str, str], cite) -> None:
    head = doc.add_paragraph()
    run = head.add_run(f"Concept {number}")
    run.bold, run.font.color.rgb = True, INDIGO
    tags = [DIFFICULTY_LABELS.get(c["difficulty"], c["difficulty"])]
    if c.get("teacher_emphasis"):
        tags.append("Teacher emphasis")
    if c.get("likely_on_test"):
        tags.append("Likely on the test")
    _note(head, "   " + " · ".join(tags), GREY)
    doc.add_heading(c["title"], level=2)
    if c.get("prerequisites"):
        _note(
            doc.add_paragraph(), "Builds on: " + ", ".join(titles.get(p, p) for p in c["prerequisites"]), GREY
        )
    quality = c.get("quality") or {}
    if quality.get("needs_checking"):
        cell = _callout(doc, "FFF7ED")
        _bold_line(cell.paragraphs[0], f"⚠ {CHECK_THIS}")
        cell.add_paragraph(CHECK_THIS_BODY)
        for p in quality.get("problems", []):
            cell.add_paragraph(p, style="List Bullet")
    _blocks(doc, c["concept"])
    cite(doc.paragraphs[-1], c.get("source_refs"))
    if c.get("levels"):
        doc.add_heading("Explain it three ways", level=3)
        for key, label in LEVEL_TABS:
            _bold_line(doc.add_paragraph(), label)
            _blocks(doc, c["levels"].get(key, ""))
    doc.add_heading("Why it matters", level=3)
    _blocks(doc, c["why_it_matters"])
    if c.get("how_it_works"):
        doc.add_heading("How it works", level=3)
        for i, step in enumerate(c["how_it_works"], start=1):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.25)
            _bold_line(p, f"{i}. ")
            _inline(p, step)
    diagram = c.get("diagram")
    if diagram and diagram.get("type") != "none" and diagram.get("mermaid"):
        _diagram(doc, diagram)
    if c.get("examples"):
        doc.add_heading("Examples" if len(c["examples"]) > 1 else "Example", level=3)
        for ex in c["examples"]:
            _blocks(doc, ex["text"])
            label = _indented(doc)
            if ex["origin"] == "source":
                _note(label, FROM_SOURCE, GREEN)
                cite(label, ex.get("source_refs"))
            else:
                _note(label, "✦ " + AI_EXAMPLE, GREY)
    if c.get("analogy"):
        doc.add_heading("Think of it like this", level=3)
        _blocks(doc, c["analogy"]["text"])
        if c["analogy"]["origin"] == "ai_generated":
            _note(_indented(doc), "✦ " + AI_ANALOGY, GREY)
    if c.get("common_mistake"):
        doc.add_heading("Common mistake", level=3)
        cell = _callout(doc, "FFFBEB")
        _inline(cell.paragraphs[0], c["common_mistake"])
    if c.get("quick_check"):
        doc.add_heading("Quick check", level=3)
        for qc in c["quick_check"]:
            p = doc.add_paragraph(style="List Bullet")
            _bold_line(p, "Q: ")
            _inline(p, qc["q"])
            a = _indented(doc)
            _bold_line(a, "Answer: ")
            _inline(a, qc["a"])


def _diagram(doc, diagram: dict) -> None:
    import pymupdf

    try:
        parsed = mermaid.parse(diagram["mermaid"])
        picture = pymupdf.open(stream=svg.render(parsed).encode(), filetype="svg")
        pixmap = picture[0].get_pixmap(dpi=200)
        png = pixmap.tobytes("png")
    except (mermaid.MermaidError, RuntimeError, ValueError):
        return
    width_in = min(6.0, pixmap.width / 200)
    doc.add_picture(io.BytesIO(png), width=Inches(width_in))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    if diagram.get("caption"):
        caption = doc.add_paragraph()
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _note(caption, diagram["caption"], GREY)


# ---------- Markdown to Word ----------


def _blocks(doc, text: str | None) -> None:
    """Block Markdown (paragraphs, lists, quotes, code) as Word paragraphs."""
    tokens = _md.parse(math_to_text(text or ""))
    lists: list[str] = []
    numbers: list[int] = []
    quote = 0
    for i, token in enumerate(tokens):
        if token.type == "bullet_list_open":
            lists.append("bullet")
            numbers.append(0)
        elif token.type == "ordered_list_open":
            lists.append("ordered")
            numbers.append(int(token.attrGet("start") or 1) - 1)
        elif token.type in ("bullet_list_close", "ordered_list_close"):
            lists.pop()
            numbers.pop()
        elif token.type == "blockquote_open":
            quote += 1
        elif token.type == "blockquote_close":
            quote -= 1
        elif token.type == "inline" and tokens[i - 1].type in (
            "paragraph_open",
            "heading_open",
            "th_open",
            "td_open",
        ):
            p = doc.add_paragraph()
            if lists:
                p.paragraph_format.left_indent = Inches(0.3 * len(lists))
                if i >= 2 and tokens[i - 2].type == "list_item_open":
                    if lists[-1] == "bullet":
                        p.add_run("•  ")
                    else:
                        numbers[-1] += 1
                        p.add_run(f"{numbers[-1]}.  ")
            if quote:
                p.paragraph_format.left_indent = Inches(0.4)
            heading = tokens[i - 1].type == "heading_open"
            _inline(p, token.content, bold=heading, italic=bool(quote))
        elif token.type in ("fence", "code_block"):
            p = doc.add_paragraph()
            run = p.add_run(token.content.rstrip("\n"))
            run.font.name = "Consolas"
            run.font.size = Pt(9)


def _inline(paragraph, text: str | None, bold: bool = False, italic: bool = False) -> None:
    """Inline Markdown (bold, italics, code, line breaks) as runs in one paragraph."""
    tokens = _md.parseInline(math_to_text(text or ""))
    children = tokens[0].children if tokens else []
    strong = em = 0
    for child in children or []:
        kind = child.type
        if kind == "strong_open":
            strong += 1
        elif kind == "strong_close":
            strong -= 1
        elif kind == "em_open":
            em += 1
        elif kind == "em_close":
            em -= 1
        elif kind == "softbreak":
            paragraph.add_run(" ")
        elif kind == "hardbreak":
            paragraph.add_run().add_break()
        elif kind == "code_inline":
            run = paragraph.add_run(child.content)
            run.font.name = "Consolas"
        elif kind in ("text", "html_inline"):
            run = paragraph.add_run(child.content)
            run.bold = bold or strong > 0 or None
            run.italic = italic or em > 0 or None


# ---------- small helpers ----------


def _setup(doc, title: str) -> None:
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    section = doc.sections[0]
    section.left_margin = section.right_margin = Inches(0.9)
    footer = section.footer.paragraphs[0]
    _note(footer, f"{title} · page ", GREY)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    doc.core_properties.title = title
    doc.core_properties.author = "HighSchoolNotes"


def _note(paragraph, text: str, color: RGBColor):
    run = paragraph.add_run(text)
    run.font.size = Pt(9)
    run.font.color.rgb = color
    return run


def _bold_line(paragraph, text: str):
    run = paragraph.add_run(text)
    run.bold = True
    return run


def _indented(doc):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.4)
    return p


def _table(doc, headers: list[str]):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, header in zip(table.rows[0].cells, headers, strict=True):
        _bold_line(cell.paragraphs[0], header)
        _shade(cell, "F1F5F9")
    return table


def _callout(doc, fill: str):
    """A one-cell shaded table: Word's closest thing to a coloured box."""
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    cell = table.rows[0].cells[0]
    _shade(cell, fill)
    doc.add_paragraph()
    return cell


def _shade(cell, fill: str) -> None:
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shading)
